# -*- coding: utf-8 -*-
"""Genera il sito statico in docs/ a partire dai case study Word.

Uso:
    python tools/build_site.py                 costruisce il sito
    python tools/build_site.py --new-preview NOME   crea una pagina anteprima non linkata
"""
import json, os, re, shutil, sys, secrets
from html import escape as esc
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from cases_data import AREAS, CASES, DNS, INVENTORY

TOOLS = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(TOOLS)
DOCS = os.path.join(ROOT, "docs")
CFG = json.load(open(os.path.join(TOOLS, "site.config.json"), encoding="utf-8"))
PROFILE = json.load(open(os.path.join(TOOLS, "profile.json"), encoding="utf-8"))
PROJ = json.load(open(os.path.join(TOOLS, "projects.json"), encoding="utf-8"))["projects"]
SRC = os.path.normpath(os.path.join(TOOLS, CFG["source_dir"]))
LANGS = ("it", "en")
import hashlib, time
VER = hashlib.sha256(str(time.time()).encode()).hexdigest()[:8]

UI = {
    "en": dict(
        home="Home", kicker="Portfolio",
        lead=("I am a Computer Science student at the University of Salerno. This is where I collect the software projects I have built, "
              "an independent research project and the case studies from the Google Cybersecurity Certificate."),
        cyber_kicker="Cybersecurity", cyber_title="Cybersecurity case studies",
        cyber_lead=("These case studies were completed during the Google Cybersecurity Certificate. Each one is based on a simulated "
                    "scenario provided by the course, and no real systems were analyzed. The topics include network traffic analysis, incident response, "
                    "risk and vulnerability assessment, access control, regulatory compliance and basic system hardening."),
        proj_open="Explore the project",
        nav_about="About", nav_cases="Cybersecurity", nav_projects="Projects", nav_path="Training",
        btn_cases="View the case studies", btn_about="About me", cert_kicker="Training", cert_more="Learn more about the program",
        about_kicker="Profile", about_title="About me", skills="Skills", links="Links",
        cases_kicker="Selected work", cases_title="Case studies",
        cases_note="Each case study sits under the NIST Cybersecurity Framework 2.0 function that describes it best. The filters also show the ones that touch that function.", also="Also touches", ref="Reference", recover_note="No case study is only about recovery. These two cover it in their own section, as part of the incident response.",
        n_cs="case studies", n_ar="CSF functions", n_lg="languages", all="All", read="Read the case study",
        crumb="All case studies", pdf="Download as PDF", prev="Previous", next="Next",
        foot1="Personal portfolio. The cybersecurity case studies are educational and based on simulated scenarios from the Google Cybersecurity Certificate: no live systems were tested and no real data is used.",
        foot2="They are not a formal attestation, certification or legal assessment of compliance.",
        projects="Projects", projects_note="Software projects, with code and instructions on GitHub.",
        previews="Previews", previews_note="Work in progress. These pages are not listed publicly.",
        open="Open", skip="Skip to content", desc_home="Portfolio of a Computer Science student: software projects, independent research and cybersecurity case studies.",
        credential="View the credential", source="Source", courses_n="%d courses",
        proj_kicker="Software", proj_title="Projects",
        proj_lead=("Software projects developed independently or as part of my university studies. The source code and instructions for running each project are available on GitHub."),
        proj_empty="Projects will be published here.", proj_code="View the code on GitHub", proj_how="How to run it", proj_results="Key findings",
        proj_all="View the projects", proj_teaser="Software projects, with code and instructions on GitHub.",
        notfound="Page not found", back="Back to the home page"),
    "it": dict(
        home="Home", kicker="Portfolio",
        lead=("Sono studente di Informatica all'Università degli Studi di Salerno. Qui raccolgo i progetti software che ho sviluppato, "
              "un lavoro di ricerca indipendente e i case study del Google Cybersecurity Certificate."),
        cyber_kicker="Cybersecurity", cyber_title="Case study di cybersecurity",
        cyber_lead=("I case study sono stati realizzati durante il Google Cybersecurity Certificate. Ogni lavoro si basa su uno scenario simulato "
                    "proposto dal corso: non sono stati analizzati sistemi reali. Gli argomenti comprendono l'analisi del traffico di rete, la risposta agli incidenti, "
                    "la valutazione di rischi e vulnerabilità, il controllo degli accessi, la conformità normativa e l'hardening di base dei sistemi."),
        proj_open="Scopri il progetto",
        nav_about="Profilo", nav_cases="Cybersecurity", nav_projects="Progetti", nav_path="Formazione",
        btn_cases="Vai ai case study", btn_about="Profilo", cert_kicker="Formazione", cert_more="Approfondisci il percorso",
        about_kicker="Profilo", about_title="Riguardo me", skills="Competenze", links="Link",
        cases_kicker="Lavori selezionati", cases_title="Case study",
        cases_note="Ogni case study sta sotto la funzione del NIST Cybersecurity Framework 2.0 che lo descrive meglio. I filtri mostrano anche quelli che la toccano.", also="Tocca anche", ref="Riferimento", recover_note="Nessun case study riguarda solo il ripristino. Questi due lo trattano in una sezione dedicata, dentro la risposta all'incidente.",
        n_cs="case study", n_ar="funzioni CSF", n_lg="lingue", all="Tutti", read="Leggi il case study",
        crumb="Tutti i case study", pdf="Scarica il PDF", prev="Precedente", next="Successivo",
        foot1="Portfolio personale. I case study di cybersecurity sono didattici e basati su scenari simulati del Google Cybersecurity Certificate: non sono stati testati sistemi reali e non si usano dati reali.",
        foot2="Non costituiscono un'attestazione formale, una certificazione né una valutazione legale di conformità.",
        projects="Progetti", projects_note="Progetti software, con codice e istruzioni su GitHub.",
        previews="Anteprime", previews_note="Lavori in corso. Queste pagine non sono elencate pubblicamente.",
        open="Apri", skip="Vai al contenuto", desc_home="Portfolio di uno studente di Informatica: progetti software, ricerca indipendente e case study di cybersecurity.",
        credential="Verifica il certificato", source="Fonte", courses_n="%d corsi",
        proj_kicker="Software", proj_title="Progetti",
        proj_lead=("Progetti software sviluppati in autonomia o nell'ambito degli studi universitari. Il codice sorgente e le istruzioni per eseguire ciascun progetto sono disponibili su GitHub."),
        proj_empty="I progetti saranno pubblicati qui.", proj_code="Vedi il codice su GitHub", proj_how="Come eseguirlo", proj_results="Risultati principali",
        proj_all="Vedi i progetti", proj_teaser="Progetti software, con codice e istruzioni su GitHub.",
        notfound="Pagina non trovata", back="Torna alla home"),
}

CSP = ("default-src 'none'; img-src 'self' data:; style-src 'self'; script-src 'self'; "
       "base-uri 'none'; form-action 'none'")

BADGE = {
    "ok": {"Implemented", "In atto", "Maintain", "Mantenere"},
    "warn": {"Partially implemented", "Parzialmente in atto", "High", "Alta"},
    "bad": {"Not implemented", "Non in atto", "Critical", "Critica"},
}


# ---------------------------------------------------------------- pagina base
ICONS = {
    "net": '<circle cx="5" cy="12" r="2.2"/><circle cx="19" cy="6" r="2.2"/><circle cx="19" cy="18" r="2.2"/><path d="M7 11l10-4M7 13l10 4"/>',
    "incident": '<path d="M12 3.5l9.5 16.5h-19z"/><path d="M12 10v4.5M12 17.4v.1"/>',
    "risk": '<path d="M3 20h18"/><path d="M6 20v-7M11 20V5M16 20v-10"/>',
    "access": '<circle cx="8" cy="14" r="4"/><path d="M11 11l9-8M16 6l3 3M14 8l2 2"/>',
    "compliance": '<rect x="5" y="4" width="14" height="17" rx="1.5"/><path d="M9 4V3h6v1M8.5 13l2.5 2.5 4.5-4.5"/>',
    "systems": '<rect x="3" y="4" width="18" height="16" rx="1.5"/><path d="M7 9l3 3-3 3M12.5 15H17"/>',
    "govern": '<rect x="5" y="4" width="14" height="17" rx="1.5"/><path d="M9 4V3h6v1M8.5 13l2.5 2.5 4.5-4.5"/>',
    "identify": '<path d="M3 20h18"/><path d="M6 20v-7M11 20V5M16 20v-10"/>',
    "protect": '<path d="M12 3l8 3v6c0 4.500-3.400 8-8 9-4.600-1-8-4.500-8-9V6z"/><path d="M8.500 12l2.500 2.500 4.500-4.500"/>',
    "detect": '<circle cx="11" cy="11" r="6.500"/><path d="M16 16l5 5M8.500 11h5"/>',
    "respond": '<path d="M12 3.500l9.500 16.500h-19z"/><path d="M12 10v4.500M12 17.400v.1"/>',
    "recover": '<path d="M4 12a8 8 0 1 0 2.600-5.900"/><path d="M4 4v4.500h4.500"/>',
}


def icon(key):
    return ('<svg class="ico" viewBox="0 0 24 24" width="26" height="26" fill="none" stroke="currentColor" stroke-width="1.5" '
            'stroke-linecap="round" stroke-linejoin="round" aria-hidden="true" focusable="false">%s</svg>' % ICONS[key])


PHOTO = None  # (larghezza, altezza) se esiste una foto


def prepare_photo():
    """Cerca tools/photo/profile.*, la riduce e la salva senza metadati (EXIF, posizione GPS)."""
    global PHOTO
    PHOTO = None
    folder = os.path.join(TOOLS, "photo")
    if not os.path.isdir(folder):
        return
    for name in sorted(os.listdir(folder)):
        if name.lower().startswith("profile.") and name.lower().rsplit(".", 1)[-1] in ("jpg", "jpeg", "png", "webp"):
            try:
                from PIL import Image, ImageOps
            except ImportError:
                print("Per usare la foto serve Pillow: pip install pillow")
                return
            im = ImageOps.exif_transpose(Image.open(os.path.join(folder, name))).convert("RGB")
            im.thumbnail((1200, 1500))
            os.makedirs(os.path.join(DOCS, "assets"), exist_ok=True)
            im.save(os.path.join(DOCS, "assets", "profile.jpg"), "JPEG", quality=88, optimize=True)  # nessun EXIF
            PHOTO = im.size
            print("Foto inclusa (%dx%d, metadati rimossi)" % im.size)
            return


def initials():
    parts = [w for w in CFG["name"].split() if w]
    return (parts[0][0] + parts[-1][0]).upper() if len(parts) > 1 else parts[0][:2].upper()


def portrait_html(lang, cls=""):
    cap = '<figcaption><b>%s</b><span>%s</span></figcaption>' % (esc(CFG["name"]), esc(PROFILE["role"][lang]))
    if PHOTO:
        return ('<figure class="portrait %s"><div class="frame"><img src="../assets/profile.jpg?v=%s" width="%d" height="%d" alt="%s"></div>%s</figure>'
                % (cls, VER, PHOTO[0], PHOTO[1], esc(PROFILE["photo_alt"][lang]), cap))
    return '<figure class="portrait placeholder %s"><div class="frame"><div class="mono" aria-hidden="true">%s</div></div>%s</figure>' % (cls, esc(initials()), cap)


def nav_html(lang, here):
    u = UI[lang]
    items = [("about.html", u["nav_about"], "about"), ("projects.html", u["nav_projects"], "projects"),
             ("cybersecurity.html", u["nav_cases"], "cyber")]
    out = []
    for href, label, key in items:
        cur = ' aria-current="page"' if key == here else ""
        out.append('<a href="%s"%s>%s</a>' % (href, cur, esc(label)))
    return '<nav class="menu" aria-label="Menu">%s</nav>' % "".join(out)


def page(lang, title, desc, body, depth, alt=None, noindex=False, here="home"):
    up = "../" * depth
    other = "en" if lang == "it" else "it"
    sw = ""
    if alt is not None:
        sw = ('<nav class="lang" aria-label="Language"><span aria-current="true">%s</span><a href="%s" hreflang="%s" lang="%s">%s</a></nav>'
              % (lang.upper(), esc(alt), other, other, other.upper()))
    robots = '<meta name="robots" content="noindex, nofollow">' if noindex else ""
    L = PROFILE["links"]
    return f"""<!doctype html>
<html lang="{lang}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="{CSP}">
<meta name="referrer" content="no-referrer">
{robots}
<title>{esc(title)}</title>
<meta name="description" content="{esc(desc)}">
<meta property="og:title" content="{esc(title)}">
<meta property="og:description" content="{esc(desc)}">
<link rel="icon" href="{up}assets/favicon.svg" type="image/svg+xml">
<link rel="stylesheet" href="{up}assets/style.css?v={VER}">
</head>
<body>
<a class="skip" href="#main">{esc(UI[lang]['skip'])}</a>
<header class="topbar"><div class="wrap bar">
{nav_html(lang, here)}
<a class="brand" href="index.html">{esc(CFG['name'])} <small>{esc(PROFILE['role'][lang])}</small></a>
{sw}
</div></header>
<main id="main">
{body}
</main>
<footer class="site-foot"><div class="wrap foot-grid">
<div><p class="foot-name">{esc(CFG['name'])}</p><p>{esc(UI[lang]['foot1'])}</p><p>{esc(UI[lang]['foot2'])}</p></div>
<ul class="foot-links"><li><a href="{esc(L['github'])}" rel="noopener noreferrer">GitHub</a></li><li><a href="{esc(L['linkedin'])}" rel="noopener noreferrer">LinkedIn</a></li></ul>
</div></footer>
<script src="{up}assets/site.js?v={VER}" defer></script>
</body>
</html>
"""


# ---------------------------------------------------------------- docx -> html
def is_bold(p):
    rs = [r for r in p.runs if r.text.strip()]
    return bool(rs) and all(r.bold for r in rs)


def is_italic(p):
    rs = [r for r in p.runs if r.text.strip()]
    return bool(rs) and all(r.italic for r in rs)


def max_size(p):
    sizes = [r.font.size.pt for r in p.runs if r.font.size and r.text.strip()]
    return max(sizes) if sizes else 0


def is_mono(p):
    return any((r.font.name or "").lower() == "consolas" for r in p.runs if r.text.strip())


def has_num(p):
    pPr = p._p.pPr
    return pPr is not None and pPr.numPr is not None


def cell_html(cell, extra=""):
    """Contenuto di una cella: istruzioni del modulo in grigio, risposta in nero."""
    paras = [p for p in cell.paragraphs if p.text.strip()]
    kinds = []
    for p in paras:
        t = p.text.strip()
        if has_num(p):
            kinds.append("b")
        elif t.startswith(("Objective:", "Obiettivo:")) or is_italic(p):
            kinds.append("q")
        else:
            kinds.append("p")
    if "b" in kinds:
        first_b = kinds.index("b")
        last_b = len(kinds) - 1 - kinds[::-1].index("b")
        for i in range(first_b):
            if kinds[i] == "p":
                kinds[i] = "q"
        # i punti elenco tra istruzione e risposta sono domande guida
        if any(k == "p" for k in kinds[last_b + 1:]):
            kinds = ["g" if k == "b" else k for k in kinds]
    elif len(paras) >= 2 and paras[0].text.strip().endswith("?"):
        kinds[0] = "q"
    out, group = [], []

    def flush():
        if group:
            cls = ' class="qlist"' if group[0][0] == "g" else ""
            out.append("<ul%s>" % cls + "".join("<li>%s</li>" % esc(t) for _, t in group) + "</ul>")
            group.clear()
    for k, p in zip(kinds, paras):
        t = p.text.strip()
        if k in ("b", "g"):
            group.append((k, t))
            continue
        flush()
        out.append('<p class="q">%s</p>' % esc(t) if k == "q" else "<p>%s</p>" % esc(t))
    flush()
    return "".join(out)


def generic_table(tbl, badges=False):
    rows = list(tbl.rows)
    if not rows:
        return ""
    ncols = len(tbl.columns)
    first_cells = rows[0].cells
    header = ncols >= 3 and all(c.text.strip() and all(r.bold for r in c.paragraphs[0].runs if r.text.strip()) for c in first_cells if c.text.strip())
    html = ['<div class="tablewrap"><table>']
    for ri, row in enumerate(rows):
        html.append("<tr>")
        for tc in row._tr.tc_lst:
            from docx.table import _Cell
            cell = _Cell(tc, tbl)
            span = tc.grid_span or 1
            cs = ' colspan="%d"' % span if span > 1 else ""
            txt = cell.text.strip()
            tag, cls = "td", ""
            if ri == 0 and header:
                tag = "th"
            elif cell is not None and tc is row._tr.tc_lst[0] and ncols <= 4:
                paras = [p for p in cell.paragraphs if p.text.strip()]
                if paras and all(is_bold(p) for p in paras) and len(row._tr.tc_lst) > 1:
                    cls = ' class="label"'
            inner = cell_html(cell)
            if badges and tag == "td":
                for kind, words in BADGE.items():
                    if txt in words:
                        inner = '<span class="badge %s">%s</span>' % (kind, esc(txt))
            html.append("<%s%s%s>%s</%s>" % (tag, cls, cs, inner, tag))
        html.append("</tr>")
    html.append("</table></div>")
    return "".join(html)


SEC_RE = re.compile(r"^(Section|Sezione|Part|Parte)\s+(\d+)\s*:\s*(.*)$", re.I)


def step_heading(label, n, title):
    title = title.strip()
    title = title[:1].upper() + title[1:]
    return '<h2 class="step"><span class="num">%s %s</span>%s</h2>' % (esc(label.capitalize()), esc(n), esc(title))


def iter_paragraphs(doc):
    """Paragrafi e tabelle in ordine. Le tabelle a colonna unica (contenitori di sezioni) vengono spianate."""
    from docx.table import _Cell
    for ch in doc.element.body.iterchildren():
        tag = ch.tag.split("}")[1]
        if tag == "p":
            yield ("p", Paragraph(ch, doc), False)
        elif tag == "tbl":
            t = Table(ch, doc)
            if t.rows and all(len(r._tr.tc_lst) == 1 for r in t.rows):
                for r in t.rows:
                    for p in _Cell(r._tr.tc_lst[0], t).paragraphs:
                        yield ("p", p, True)
            else:
                yield ("t", t, False)


def render_generic(doc):
    items = list(iter_paragraphs(doc))
    out, bullets, code, numbered = [], [], [], []
    seen_title = False

    def flush():
        if bullets:
            out.append("<ul>" + "".join("<li>%s</li>" % esc(b) for b in bullets) + "</ul>")
            bullets.clear()
        if code:
            out.append("<pre>%s</pre>" % esc("\n".join(code)))
            code.clear()
        if numbered:
            out.append('<ol class="steps">' + "".join("<li>%s</li>" % esc(n) for n in numbered) + "</ol>")
            numbered.clear()
    i = 0
    while i < len(items):
        kind, obj, in_cell = items[i]
        i += 1
        if kind == "t":
            flush()
            out.append(generic_table(obj))
            continue
        p = obj
        t = p.text.strip()
        if not t:
            continue
        style = p.style.name if p.style is not None else ""
        if is_mono(p):
            if bullets or numbered:
                flush()
            code.append(t)
            continue
        if code:
            flush()
        m = re.match(r"^(\d+)\.\s+(.*)$", t)
        if m and not has_num(p):
            if bullets or code:
                flush()
            numbered.append(m.group(2))
            continue
        if numbered:
            flush()
        if has_num(p):
            bullets.append(t)
            continue
        flush()
        sec = SEC_RE.match(t) if (is_bold(p) or style.startswith("Heading")) else None
        if sec:
            title = sec.group(3)
            # il titolo può continuare nel paragrafo successivo (riga spezzata)
            if (i < len(items) and items[i][0] == "p" and is_bold(items[i][1])
                    and not SEC_RE.match(items[i][1].text.strip()) and items[i][1].text.strip()[:1].islower()):
                title += " " + items[i][1].text.strip()
                i += 1
            out.append(step_heading(sec.group(1), sec.group(2), title))
            continue
        if style in ("Title", "Heading 1", "Heading 2", "Heading 3") or (is_bold(p) and max_size(p) >= 18):
            if not seen_title:
                seen_title = True
                continue
            out.append("<h3>%s</h3>" % esc(t) if style == "Heading 3" else "<h2>%s</h2>" % esc(t))
        elif style == "Subtitle":
            out.append('<p class="note">%s</p>' % esc(t))
        elif is_bold(p) and max_size(p) >= 12.5 and len(t) < 120:
            out.append("<h2>%s</h2>" % esc(t))
        elif is_italic(p) or t.endswith(":"):
            out.append('<p class="q">%s</p>' % esc(t))
        elif in_cell and ": " in t[:120] and not t[0].isdigit():
            lead, rest = t.split(": ", 1)
            out.append('<p><span class="lead-in">%s:</span> %s</p>' % (esc(lead), esc(rest)))
        else:
            out.append("<p>%s</p>" % esc(t))
    flush()
    return "\n".join(out)


def render_botium(doc):
    out, bullets = [], []
    cover = True
    stats_done = False

    def flush():
        if bullets:
            out.append("<ul>" + "".join("<li>%s</li>" % esc(b) for b in bullets) + "</ul>")
            bullets.clear()

    def statrow(tbl):
        cards = []
        for tc in tbl.rows[0]._tr.tc_lst:
            from docx.table import _Cell
            ps = [p.text.strip() for p in _Cell(tc, tbl).paragraphs if p.text.strip()]
            if not ps:
                continue
            big = esc(ps[0]); lab = esc(ps[1]) if len(ps) > 1 else ""; sub = esc(ps[2]) if len(ps) > 2 else ""
            cards.append('<div class="stat"><b>%s</b><span>%s</span><small>%s</small></div>' % (big, lab, sub))
        return '<div class="statrow">%s</div>' % "".join(cards)

    def is_stat(tbl):
        if len(tbl.rows) != 1 or len(tbl.columns) < 2:
            return False
        c0 = [p.text.strip() for p in tbl.rows[0].cells[0].paragraphs if p.text.strip()]
        return len(c0) >= 2 and len(c0[0]) <= 10

    for ch in doc.element.body.iterchildren():
        tag = ch.tag.split("}")[1]
        if tag == "tbl":
            tbl = Table(ch, doc)
            flush()
            if is_stat(tbl):
                out.append(statrow(tbl))
            elif len(tbl.rows) == 1 and len(tbl.columns) == 1:
                ps = [p.text.strip() for p in tbl.rows[0].cells[0].paragraphs if p.text.strip()]
                if len(ps) >= 2:
                    out.append('<div class="callout"><b>%s</b><p>%s</p></div>' % (esc(ps[0]), esc(" ".join(ps[1:]))))
            else:
                out.append(generic_table(tbl, badges=True))
            continue
        if tag != "p":
            continue
        p = Paragraph(ch, doc)
        t = p.text.strip()
        xml = ch.xml
        if cover:
            if 'w:type="page"' in xml:
                cover = False
            elif len(p.runs) == 2 and p.runs[0].text.strip() in ("Focus", "Ambito"):
                out.append('<p class="focus"><b>%s</b>%s</p>' % (esc(p.runs[0].text.strip()), esc(p.runs[-1].text.strip())))
            continue
        if not t or set(t) <= set("━ "):
            continue
        runs = [r for r in p.runs if r.text.strip()]
        if len(p.runs) >= 2 and p.runs[0].text.strip().isdigit() and len(p.runs[0].text.strip()) <= 2:
            flush()
            title = "".join(r.text for r in p.runs[1:]).strip()
            out.append('<h2><span class="num">%s</span>%s</h2>' % (esc(p.runs[0].text.strip()), esc(title)))
        elif t.startswith("•"):
            bullets.append(t.lstrip("•").strip())
        else:
            flush()
            sz = max_size(p)
            if len(runs) == 1 and runs[0].bold and sz >= 11:
                out.append("<h3>%s</h3>" % esc(t))
            elif sz and sz <= 8.5 and not runs[0].bold:
                out.append('<p class="note">%s</p>' % esc(t))
            else:
                out.append("<p>%s</p>" % esc(t))
    flush()
    return "\n".join(out)


def render_dns(lang):
    d = DNS[lang]
    short = {"it": ["Riepilogo del problema", "Analisi dei dati e causa probabile"],
             "en": ["Summary of the problem", "Analysis and likely cause"]}[lang]
    out = []
    for idx, (label, q, boxes) in enumerate(d["parts"]):
        n = re.sub(r"\D", "", label)
        word = re.sub(r"\d+", "", label).strip().capitalize()
        out.append('<h2 class="step"><span class="num">%s %s</span>%s</h2><p class="q">%s</p><div class="boxes">' % (esc(word), esc(n), esc(short[idx]), esc(q)))
        for head, paras in boxes:
            out.append('<div class="box">%s%s</div>' % ("<h3>%s</h3>" % esc(head) if head else "", "".join("<p>%s</p>" % esc(x) for x in paras)))
        out.append("</div>")
    return "\n".join(out)


def render_inventory(lang):
    d = INVENTORY[lang]

    def table(h, rows):
        t = ['<div class="tablewrap"><table><tr>%s</tr>' % "".join("<th>%s</th>" % esc(x) for x in h)]
        for r in rows:
            t.append("<tr>%s</tr>" % "".join("<td>%s</td>" % esc(x) for x in r))
        t.append("</table></div>")
        return "".join(t)
    return "<h2>%s</h2>%s<h2>%s</h2>%s" % (esc(d["t1"]), table(d["h1"], d["rows"]), esc(d["t2"]), table(d["h2"], d["cats"]))


def article_body(case, lang):
    kind = case["kind"]
    if kind == "dns":
        return render_dns(lang)
    if kind == "inventory":
        return render_inventory(lang)
    path = os.path.join(SRC, lang.upper(), case["stem"] + ".docx")
    doc = Document(path)
    return render_botium(doc) if kind == "botium" else render_generic(doc)


# ---------------------------------------------------------------- pagine
def active_projects(kind):
    flag = "show_public_projects" if kind == "public" else "show_previews"
    if not CFG.get(flag):
        return []
    return [p for p in PROJ if p.get("visibility") == kind and p.get("enabled")]


def projects_html(lang):
    """Schede dei progetti nella home: compaiono solo se ci sono progetti pubblici attivi."""
    u = UI[lang]
    pub = active_projects("public")
    if not pub:
        return ""
    cards = []
    for p in pub:
        tone = p.get("tone", "risk")
        im = p.get("image")
        thumb = ('<img class="card-img" src="../assets/img/%s" width="%d" height="%d" alt="%s" loading="lazy">'
                 % (esc(im["file"]), im["w"], im["h"], esc(im["alt"][lang]))) if im else ""
        cards.append('<li class="card a-%s"><a href="projects.html#%s">%s<span class="card-top">%s<span class="tag">%s</span></span>'
                     '<h3>%s</h3><p>%s</p><span class="skills">%s</span><span class="go">%s &rarr;</span></a></li>'
                     % (tone, esc(p["id"]), thumb, icon(tone) if tone in ICONS else "", esc(p["status"][lang]), esc(p["title"]),
                        esc((p.get("teaser") or p["summary"])[lang]), esc(p["tags"][lang]), esc(u["proj_open"])))
    return ('<section class="cases" id="progetti"><div class="wrap"><p class="kicker center">%s</p><h2 class="section-title">%s</h2>'
            '<p class="section-note">%s</p><ul class="grid fit">%s</ul></div></section>'
            % (esc(u["proj_kicker"]), esc(u["proj_title"]), esc(u["proj_teaser"]), "".join(cards)))


def project_block(p, lang, u, preview=False):
    how = ""
    if p.get("how"):
        how = '<h3>%s</h3><ol class="steps">%s</ol>' % (esc(u["proj_how"]), "".join("<li>%s</li>" % esc(x) for x in p["how"][lang]))
    results = ""
    if p.get("highlights"):
        results = '<h3>%s</h3><ul>%s</ul>' % (esc(u["proj_results"]), "".join("<li>%s</li>" % esc(x) for x in p["highlights"][lang]))
    note = '<p class="note">%s</p>' % esc(p["note"][lang]) if p.get("note") else ""
    if preview:
        href, label, rel = "../preview/%s/" % p["token"], u["open"], ""
    else:
        href, label, rel = p["url"], u["proj_code"], ' rel="noopener noreferrer"'
    tags = "".join("<li>%s</li>" % esc(t.strip()) for t in p["tags"][lang].split(","))
    shot = ""
    im = p.get("image")
    if im and not preview:
        file, w, h = (im["detail"], im["dw"], im["dh"]) if im.get("detail") else (im["file"], im["w"], im["h"])
        shot = ('<figure class="project-shot"><img src="../assets/img/%s" width="%d" height="%d" alt="%s" loading="lazy">'
                '<figcaption>%s</figcaption></figure>' % (esc(file), w, h, esc(im["alt"][lang]), esc(im["caption"][lang])))
    return (f'<article class="project a-{p.get("tone", "risk")}" id="{esc(p["id"])}">'
            f'<div class="project-meta"><span class="tag">{esc(p["status"][lang])}</span><h2>{esc(p["title"])}</h2>'
            f'<ul class="chips">{tags}</ul><a class="btn solid" href="{esc(href)}"{rel}>{esc(label)}</a></div>'
            f'<div class="project-body">{shot}<p class="project-sum">{esc(p["summary"][lang])}</p>{results}{how}{note}</div></article>')


def build_projects(lang):
    u = UI[lang]
    pub = active_projects("public")
    prev = active_projects("preview")
    blocks = "".join(project_block(p, lang, u) for p in pub)
    if prev:
        blocks += '<h2 class="section-title sub">%s</h2><p class="section-note">%s</p>' % (esc(u["previews"]), esc(u["previews_note"]))
        blocks += "".join(project_block(p, lang, u, preview=True) for p in prev)
    if not blocks:
        blocks = '<p class="lead center">%s</p>' % esc(u["proj_empty"])
    body = f"""<div class="wrap">
<header class="about-head">
<p class="kicker">{esc(u['proj_kicker'])}</p>
<h1>{esc(u['proj_title'])}</h1>
<hr class="rule">
<p class="lead center">{esc(u['proj_lead'])}</p>
</header>
<div class="projects">{blocks}</div>
</div>"""
    other = "en" if lang == "it" else "it"
    return page(lang, f"{u['proj_title']} | {CFG['name']}", u["proj_lead"], body, 1, alt="../%s/projects.html" % other, here="projects")


def cert_band(lang, more_href, more_label):
    """Fascia blu sul Google Cybersecurity Certificate, con scheda dei dati e link alla verifica."""
    u = UI[lang]
    C = PROFILE["certificate"]
    facts = "".join("<div><dt>%s</dt><dd>%s</dd></div>" % (esc(f["label"][lang]), esc(f["value"][lang])) for f in C["facts"])
    cert_link = ""
    if C.get("credential_url"):
        cert_link = '<a class="btn light" href="%s" rel="noopener noreferrer">%s</a>' % (esc(C["credential_url"]), esc(u["credential"]))
    return f"""<section class="cert-band" id="formazione">
<div class="wrap cert-grid">
<div>
<p class="kicker">{esc(u['cert_kicker'])}</p>
<h2>{esc(C['title'][lang])}</h2>
<p>{esc(C['band'][lang])}</p>
<div class="actions"><a class="btn light" href="{more_href}">{esc(more_label)}</a>{cert_link}</div>
</div>
<dl class="facts">{facts}</dl>
</div>
</section>"""


def build_home(lang):
    u = UI[lang]
    body = f"""<section class="hero">
<div class="wrap hero-grid">
<div class="hero-text">
<p class="kicker">{esc(u['kicker'])}</p>
<h1>{esc(CFG['site_title'][lang])}</h1>
<hr class="rule">
<p class="lead">{esc(u['lead'])}</p>
<div class="actions"><a class="btn solid" href="projects.html">{esc(u['proj_all'])}</a><a class="btn" href="about.html">{esc(u['btn_about'])}</a></div>
</div>
{portrait_html(lang, "hero-photo")}
</div>
</section>
{projects_html(lang)}
{cert_band(lang, "cybersecurity.html", u['btn_cases'])}"""
    other = "en" if lang == "it" else "it"
    return page(lang, f"{CFG['site_title'][lang]} | {CFG['name']}", u["desc_home"], body, 1, alt="../%s/index.html" % other)


def build_cyber(lang):
    u = UI[lang]
    chips = ['<button class="chip" data-filter="all" aria-pressed="true">%s</button>' % esc(u["all"])]
    for k, v in AREAS.items():
        chips.append('<button class="chip c-%s" data-filter="%s" aria-pressed="false">%s</button>' % (k, k, esc(v[lang])))
    cards = []
    for c in CASES:
        also_names = ", ".join(AREAS[k][lang] for k in c.get("also", []))
        also_line = '<span class="also">%s: %s</span>' % (esc(u["also"]), esc(also_names)) if also_names else ""
        cards.append(
            '<li class="card a-%s" data-area="%s" data-also="%s"><a href="%s.html"><span class="card-top">%s<span class="tag">%s</span></span>'
            '<h3>%s</h3><p>%s</p>%s<span class="skills">%s</span><span class="go">%s &rarr;</span><span class="refbadge">%s</span></a></li>'
            % (c["area"], c["area"], " ".join(c.get("also", [])), c["slug"], icon(c["area"]), esc(AREAS[c["area"]][lang]), esc(c["title"][lang]),
               esc(c["summary"][lang]), also_line, esc(c["skills"][lang]), esc(u["read"]), esc(u["ref"])))
    body = f"""<div class="wrap">
<header class="about-head">
<p class="kicker">{esc(u['cyber_kicker'])}</p>
<h1>{esc(u['cyber_title'])}</h1>
<hr class="rule">
<p class="lead center">{esc(u['cyber_lead'])}</p>
</header>
</div>
{cert_band(lang, "about.html#percorso", u['cert_more'])}
<section class="cases" id="case-study">
<div class="wrap">
<p class="kicker center">{esc(u['cases_kicker'])}</p>
<h2 class="section-title">{esc(u['cases_title'])}</h2>
<p class="section-note">{esc(u['cases_note'])}</p>
<div class="filters" role="group" aria-label="Filter">{''.join(chips)}</div>
<p class="filter-note" id="filter-note-recover" hidden>{esc(u['recover_note'])}</p>
<ul class="grid">{''.join(cards)}</ul>
</div>
</section>"""
    other = "en" if lang == "it" else "it"
    return page(lang, f"{u['cyber_title']} | {CFG['name']}", u["cyber_lead"], body, 1, alt="../%s/cybersecurity.html" % other, here="cyber")


def build_about(lang):
    u = UI[lang]
    P = PROFILE
    C = P["certificate"]
    intro = "".join("<p>%s</p>" % esc(t) for t in P["intro"][lang])
    glance = "".join("<div><dt>%s</dt><dd>%s</dd></div>" % (esc(i["label"][lang]), esc(i["value"][lang])) for i in P["side"]["items"])
    links = '<ul class="linklist"><li><a href="%s" rel="noopener noreferrer">GitHub</a></li><li><a href="%s" rel="noopener noreferrer">LinkedIn</a></li></ul>' % (
        esc(P["links"]["github"]), esc(P["links"]["linkedin"]))
    facts = "".join('<div class="fact"><dt>%s</dt><dd>%s</dd></div>' % (esc(f["label"][lang]), esc(f["value"][lang])) for f in C["facts"])
    paras = "".join("<p>%s</p>" % esc(t) for t in C["paragraphs"][lang])
    cred = ""
    if C.get("credential_url"):
        cred = '<p><a class="btn" href="%s" rel="noopener noreferrer">%s</a></p>' % (esc(C["credential_url"]), esc(u["credential"]))
    roles = ""
    if C.get("roles"):
        roles = ('<h3>%s</h3><ul class="chips">%s</ul><p class="note">%s</p>'
                 % (esc(C["roles_title"][lang]), "".join("<li>%s</li>" % esc(r) for r in C["roles"][lang]), esc(C["roles_note"][lang])))
    courses = ""
    if C.get("courses"):
        lis = "".join("<li><span><b>%s</b><br>%s</span></li>" % (esc(c["title"]), esc(c["text"][lang])) for c in C["courses"])
        courses = ('<h3>%s</h3><ol class="courses">%s</ol><p class="note">%s</p>'
                   % (esc(C["courses_title"][lang]), lis, esc(C["courses_note"][lang])))
    cert = f"""<section class="cert-panel a-net" id="percorso">
<p class="kicker">{esc(u['cert_kicker'])}</p>
<h2>{esc(C['title'][lang])}</h2>
{paras}
<dl class="factgrid">{facts}</dl>
{roles}
{courses}
<p>{esc(C['link'][lang])}</p>
{cred}
<p class="source">{esc(C['sources'][lang])} <a href="{esc(C['source_url'])}" rel="noopener noreferrer">Coursera</a></p>
</section>"""
    skills = ""
    if P.get("skills"):
        S = P["skills"]
        cols = {1: [], 2: []}
        for g in S["groups"]:
            lis = []
            for it in g["items"]:
                name = it["name"] if isinstance(it["name"], str) else it["name"][lang]
                if it.get("links"):
                    link = '<span class="skill-links">%s</span>' % "".join('<a class="pill" href="%s" aria-label="%s, %s">%s <i aria-hidden="true">&rarr;</i></a>'
                                         % (esc(l["href"]), esc(name), esc(l["label"]), esc(l["label"])) for l in it["links"])
                elif it.get("href"):
                    link = ' <a href="%s" aria-label="%s">&rarr;</a>' % (esc(it["href"]), esc(name))
                else:
                    link = ""
                lis.append("<li><b>%s</b><span>%s%s</span></li>" % (esc(name), esc(it["text"][lang]), link))
            note = '<p class="group-note">%s</p>' % esc(g["note"][lang]) if g.get("note") else ""
            cols[g.get("col", 1)].append('<div class="skill-group a-%s"><h3>%s</h3>%s<ul>%s</ul></div>'
                                         % (g["tone"], esc(g["title"][lang]), note, "".join(lis)))
        skills = ('<section class="block skills" id="skills"><h2>%s</h2><p class="skills-lead">%s</p><div class="skill-groups">%s</div></section>'
                  % (esc(S["title"][lang]), esc(S["lead"][lang]),
                     "".join('<div class="skill-col">%s</div>' % "".join(cols[k]) for k in (1, 2))))
    secs = []
    for sec in P["sections"]:
        items = []
        for it in sec["items"]:
            text = it["text"][lang]
            items.append('<li class="entry"><h3>%s</h3><p class="meta">%s</p>%s</li>'
                         % (esc(it["title"][lang]), esc(it["meta"][lang]), "<p>%s</p>" % esc(text) if text else ""))
        secs.append('<section class="block a-%s" id="%s"><h2>%s</h2><ol class="timeline">%s</ol></section>' % (sec["tone"], esc(sec["id"]), esc(sec["title"][lang]), "".join(items)))
    body = f"""<div class="wrap">
<header class="about-head">
<p class="kicker">{esc(u['about_kicker'])}</p>
<h1>{esc(u['about_title'])}</h1>
<hr class="rule">
</header>
<div class="about-grid">
<aside class="about-side">
{portrait_html(lang)}
<div class="glance"><h2>{esc(P['side']['title'][lang])}</h2><dl>{glance}</dl>{links}</div>
</aside>
<article class="article about">
<div class="intro">{intro}</div>
{skills}
{cert}
{''.join(secs)}
</article>
</div>
</div>"""
    other = "en" if lang == "it" else "it"
    return page(lang, f"{u['about_title']} | {CFG['name']}", P["role"][lang] + ". " + P["intro"][lang][0], body, 1, alt="../%s/about.html" % other, here="about")


def build_case(lang, i):
    u = UI[lang]
    c = CASES[i]
    pdf = ""
    if c["kind"] != "inventory":
        pdf = '<a class="btn" href="../pdf/%s/%s.pdf" download>%s</a>' % (lang, c["slug"], esc(u["pdf"]))
    prev_c = CASES[i - 1] if i > 0 else None
    next_c = CASES[i + 1] if i < len(CASES) - 1 else None
    pager = '<nav class="pager" aria-label="Case studies">'
    pager += ('<a class="prev" href="%s.html"><small>&larr; %s</small>%s</a>' % (prev_c["slug"], esc(u["prev"]), esc(prev_c["title"][lang]))) if prev_c else "<span></span>"
    pager += ('<a class="next" href="%s.html"><small>%s &rarr;</small>%s</a>' % (next_c["slug"], esc(u["next"]), esc(next_c["title"][lang]))) if next_c else ""
    pager += "</nav>"
    skills = "".join("<li>%s</li>" % esc(s.strip()) for s in c["skills"][lang].split(","))
    body = f"""<div class="wrap">
<header class="article-head a-{c['area']}">
<p class="crumbs"><a href="cybersecurity.html">&larr; {esc(u['crumb'])}</a></p>
<span class="tag a-{c['area']}">{esc(AREAS[c['area']][lang])}</span>{('<span class="also-line">%s: %s</span>' % (esc(u['also']), esc(", ".join(AREAS[k][lang] for k in c.get('also', []))))) if c.get('also') else ''}
<h1>{esc(c['title'][lang])}</h1>
<hr class="rule">
<p class="lead">{esc(c['summary'][lang])}</p>
<ul class="skill-list">{skills}</ul>
{pdf}
</header>
<article class="article a-{c['area']}">
{article_body(c, lang)}
{pager}
</article>
</div>"""
    other = "en" if lang == "it" else "it"
    return page(lang, f"{c['title'][lang]} | {CFG['name']}", c["summary"][lang], body, 1, alt="../%s/%s.html" % (other, c["slug"]), here="cyber")


# ---------------------------------------------------------------- scrittura
def write(path, text):
    if path.endswith(".html"):
        text = text.replace("—", ", ")  # niente trattini lunghi usati come pausa
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)


def build():
    # pulizia delle sole parti generate: docs/preview non si tocca mai
    for d in ("it", "en", "assets", "pdf"):
        shutil.rmtree(os.path.join(DOCS, d), ignore_errors=True)
    os.makedirs(DOCS, exist_ok=True)
    shutil.copytree(os.path.join(TOOLS, "assets"), os.path.join(DOCS, "assets"))
    prepare_photo()
    # le pagine
    for lang in LANGS:
        write(os.path.join(DOCS, lang, "index.html"), build_home(lang))
        write(os.path.join(DOCS, lang, "about.html"), build_about(lang))
        write(os.path.join(DOCS, lang, "cybersecurity.html"), build_cyber(lang))
        write(os.path.join(DOCS, lang, "projects.html"), build_projects(lang))
        for i, c in enumerate(CASES):
            write(os.path.join(DOCS, lang, c["slug"] + ".html"), build_case(lang, i))
            if c["kind"] != "inventory":
                src = os.path.join(SRC, lang.upper(), "PDF", c["stem"] + ".pdf")
                if os.path.exists(src):
                    os.makedirs(os.path.join(DOCS, "pdf", lang), exist_ok=True)
                    shutil.copy(src, os.path.join(DOCS, "pdf", lang, c["slug"] + ".pdf"))
                else:
                    print("PDF mancante:", src)
    # radice: scelta della lingua
    write(os.path.join(DOCS, "index.html"), f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="{CSP}">
<title>{esc(CFG['name'])}</title><link rel="icon" href="assets/favicon.svg" type="image/svg+xml">
<link rel="stylesheet" href="assets/style.css?v={VER}"><script src="assets/redirect.js?v={VER}"></script></head>
<body><main class="wrap hero"><p class="kicker">Portfolio</p><h1>{esc(CFG['name'])}</h1>
<p class="lead"><a href="it/">Italiano</a> &nbsp;|&nbsp; <a href="en/">English</a></p></main></body></html>
""")
    write(os.path.join(DOCS, "404.html"), f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="{CSP}"><meta name="robots" content="noindex">
<title>404</title><link rel="stylesheet" href="{CFG.get('base_path', '')}/assets/style.css"></head>
<body><main class="wrap hero"><h1>404</h1><p class="lead">{UI['en']['notfound']} / {UI['it']['notfound']}</p></main></body></html>
""")
    write(os.path.join(DOCS, "robots.txt"), "User-agent: *\nAllow: /\nDisallow: /preview/\n")
    write(os.path.join(DOCS, ".nojekyll"), "")
    print("Sito generato in", DOCS, "(%d case study x %d lingue)" % (len(CASES), len(LANGS)))


def new_preview(name):
    token = secrets.token_urlsafe(12)
    folder = os.path.join(DOCS, "preview", token)
    write(os.path.join(folder, "index.html"), f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="{CSP}">
<meta name="robots" content="noindex, nofollow"><meta name="referrer" content="no-referrer">
<title>{esc(name)} (preview)</title><link rel="stylesheet" href="../../assets/style.css"></head>
<body><main class="wrap hero"><p class="kicker">Preview</p><h1>{esc(name)}</h1>
<p class="lead">Demo page. Use mock data only: no real data, keys or source code.</p></main></body></html>
""")
    entry = {"id": re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-"), "visibility": "preview", "enabled": False, "token": token,
             "title": name, "summary": {"it": "", "en": ""}, "status": {"it": "Anteprima", "en": "Preview"}, "tags": ""}
    p = os.path.join(TOOLS, "projects.json")
    data = json.load(open(p, encoding="utf-8"))
    data["projects"].append(entry)
    write(p, json.dumps(data, indent=2, ensure_ascii=False) + "\n")
    print("Anteprima creata: docs/preview/%s/index.html" % token)
    print("Resta nascosta finché non metti enabled=true in projects.json e show_previews=true in site.config.json.")


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "--new-preview":
        new_preview(" ".join(sys.argv[2:]))
    else:
        build()
